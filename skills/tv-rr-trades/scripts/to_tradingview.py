"""Put logged trades back on the clipboard as TradingView Long/Short Position drawings, ready to paste into a chart.

usage: to_tradingview.py [--workbook PATH] [--id ID ...] [--date YYYY-MM-DD[..YYYY-MM-DD]] [--direction long|short]
                         [--outcome TP|Stop|…] [--last N] [--min-minutes 40] [--label TEXT] [--out FILE.html] [--dry-run]

At least one selector (--id, --date, --direction, --outcome, --last) is required. Filters combine. Each drawing keeps
the trade's drawing id, so copying it back out of TradingView updates its row instead of adding a new one, and is
named "<label> <entry date and time> <direction> <outcome> <result R>" (label defaults to the workbook's name, less
"Back/Forward/Replay Testing"). Entry, stop and TP are rebuilt from the row; a closed trade gets its close point at
the exit TradingView drew (else the logged exit). Each box is at least --min-minutes wide (default 40) so
short trades stay visible; the close point stays at the real exit. The style comes from position_template.json.

Writing the clipboard needs macOS (run with --with pyobjc-framework-Cocoa). Elsewhere use --out and open the file.
"""
import argparse, copy, datetime as dt, html, json, random, re, string, sys, uuid
from pathlib import Path
from openpyxl import load_workbook
import config

TEMPLATE = Path(__file__).with_name("position_template.json")
BAR_SECONDS = {"1": 60, "3": 180, "5": 300, "15": 900, "30": 1800, "60": 3600, "240": 14400, "D": 86400}


def rows_from(path):
    ws = load_workbook(path, data_only=True, read_only=True)["Trades"]
    it = ws.iter_rows(values_only=True); hdr = next(it)
    return [dict(zip(hdr, r)) for r in it if any(v is not None for v in r)]


def default_label(path):
    return re.sub(r"\s*(Back|Forward|Replay)\s*Testing\s*$", "", Path(path).stem, flags=re.I).strip()


def epoch(naive_utc):
    return int(naive_utc.replace(tzinfo=dt.timezone.utc).timestamp())


def drawing(row, base, label, min_minutes=40):
    sym, tf = row["Symbol"], str(row.get("Timeframe") or "5")
    tick = config.symbol(sym).get("tick")
    if not tick:
        sys.exit(f"No tick size for {sym}. Set it with: config.py symbol {sym} tick <size>")
    d = copy.deepcopy(base); src = d["source"]; st = src["state"]
    long_ = str(row["Direction"]).lower() == "long"
    src["type"] = "LineToolRiskRewardLong" if long_ else "LineToolRiskRewardShort"
    src["id"] = row.get("Drawing id") or "".join(random.choices(string.ascii_letters + string.digits, k=6))
    d["modelId"] = str(uuid.uuid4())
    r = row.get("Result R")
    st.update(symbol=sym, interval=tf,
              stopLevel=round(abs(row["Entry"] - row["Stop"]) / tick),
              profitLevel=round(abs(row["TP planned"] - row["Entry"]) / tick),
              title=" ".join(x for x in (label, f'{row["Entry time"]:%-d %b %y %H:%M}', str(row["Direction"]),
                                         str(row.get("Outcome") or ""), f"{r:+g}R" if isinstance(r, (int, float)) else "") if x))
    bar = BAR_SECONDS.get(tf, 300)
    entry_t = epoch(row["Entry (UTC)"]); shown_offset = row["Entry time"] - row["Entry (UTC)"]
    exit_shown = row.get("TradingView exit") or row.get("Exit time")
    pt = lambda t, p: {"time_t": t, "offset": 0, "price": p, "interval": tf}
    pts = [pt(entry_t, row["Entry"])]
    if exit_shown and row.get("Exit price") is not None:
        exit_t = epoch(exit_shown - shown_offset)
        right = max(exit_t + 2 * bar, entry_t + min_minutes * 60)
        pts += [pt(right, row["Entry"]), pt(entry_t, row["Entry"]), pt(exit_t, row["Exit price"])]
    else:
        pts.append(pt(entry_t + max(12 * bar, min_minutes * 60), row["Entry"]))
    src["points"] = pts; d["geometry"] = d["geometry"][:len(pts)]
    return d


def to_clipboard(clip_html, text):
    from AppKit import NSPasteboard, NSPasteboardTypeHTML, NSPasteboardTypeString
    pb = NSPasteboard.generalPasteboard(); pb.clearContents()
    pb.setString_forType_(clip_html, NSPasteboardTypeHTML); pb.setString_forType_(text, NSPasteboardTypeString)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook"); ap.add_argument("--id", nargs="+", default=[])
    ap.add_argument("--date", help="YYYY-MM-DD or YYYY-MM-DD..YYYY-MM-DD, by entry date in the log's timezone")
    ap.add_argument("--direction", choices=["long", "short"]); ap.add_argument("--outcome")
    ap.add_argument("--last", type=int, help="the N most recent entries (after the other filters)")
    ap.add_argument("--min-minutes", type=int, default=40, help="minimum width of each box, so short trades stay visible")
    ap.add_argument("--label"); ap.add_argument("--out"); ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    if not (a.id or a.date or a.direction or a.outcome or a.last):
        sys.exit("Say which trades: --id, --date, --direction, --outcome and/or --last N.")
    path = config.workbook(a.workbook); rows = [r for r in rows_from(path) if r.get("Entry") is not None and r.get("Entry (UTC)")]
    if a.id: rows = [r for r in rows if r.get("Drawing id") in a.id]
    if a.date:
        lo, _, hi = a.date.partition(".."); lo, hi = dt.date.fromisoformat(lo), dt.date.fromisoformat(hi or lo)
        rows = [r for r in rows if lo <= r["Entry time"].date() <= hi]
    if a.direction: rows = [r for r in rows if str(r["Direction"]).lower() == a.direction]
    if a.outcome: rows = [r for r in rows if str(r.get("Outcome") or "").lower() == a.outcome.lower()]
    rows.sort(key=lambda r: r["Entry (UTC)"])
    if a.last: rows = rows[-a.last:]
    if not rows: sys.exit("No trades match.")
    label = a.label if a.label is not None else default_label(path)
    base = json.loads(TEMPLATE.read_text())
    sources = [drawing(r, base, label, a.min_minutes) for r in rows]
    for s in sources: print(s["source"]["state"]["title"])
    title = "Drawings" if len(sources) > 1 else f'{str(rows[0]["Direction"]).lower()} position'
    payload = html.escape(json.dumps({"sources": sources, "title": title}, separators=(",", ":")), quote=True)
    clip = f'<meta charset="utf-8"><span data-tradingview-clip="{payload}">{title}</span>'
    if a.out: Path(a.out).expanduser().write_text(clip); print(f"written to {a.out}")
    if a.dry_run: print(f"{len(sources)} drawing(s) — dry run, clipboard untouched"); return
    if not a.out or sys.platform == "darwin":
        if sys.platform != "darwin": sys.exit("The clipboard can only be written on macOS; use --out FILE.html.")
        to_clipboard(clip, title)
        print(f"{len(sources)} drawing(s) on the clipboard: click a {rows[0]['Symbol']} chart in TradingView and press ⌘V")


if __name__ == "__main__":
    main()
