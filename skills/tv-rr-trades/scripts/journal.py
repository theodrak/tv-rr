"""Journal columns for a trade log: why a trade was filtered or missed, its confluences, a grade and its reasons, notes.

usage: journal.py init WORKBOOK [--from OTHER_WORKBOOK] [--distance N --unit pts|pips]

init adds two sheets and, on the next extract.py run, the journal columns. A workbook has journal columns only when it
has a "Lists" sheet, so other logs are untouched.
  Lists            the dropdown choices: column A filter reasons, B grade reasons, C confluences (from row 2, no gaps).
                   Starts with a few generic examples, or a copy of OTHER_WORKBOOK's lists with --from. The user edits
                   them; every dropdown picks up new rows. Column E is a guide the script rewrites.
  Auto confluence  which levels enrich.py fills in by itself and how close counts: Distance and Unit (pts = price
                   units, pips = 10 ticks, as on 5-decimal FX quotes), then one row per level: Level (VWAP, SMA n,
                   EMA n, Previous day high, Previous day low), On (Yes/No) and the Name to write, which should match
                   the Confluence list. No distance = no auto-fill. Copied too with --from.
extract.py keeps both sheets and every journal value on each rebuild.

Columns:
  Decision            Taken (the default for every trade added), Filtered or Missed. Taken with filter reasons = taken
                      because something outweighed them (say what in Filter notes)
  Filter 1-3          dropdown from Lists!A
  Filter notes        free text
  Confluence 1-3      dropdown from Lists!C; auto levels are added to empty slots
  Auto confluence     what the script found (name and distance); a level the user deleted from the slots is not re-added
  Grade               A+ / A / B / C / D / F, the user's call
  Rule grade          formula: C = 2+ "-" grade reasons; B = one "-";
                      A+ = no "-" and 2+ pluses ("+" grade reasons and confluences); else A
  Grade reason 1-3    dropdown from Lists!B; items start with "+" or "-"
  Grade notes, General notes   free text
Safeguard (live in Excel, nothing is changed for the user): Decision turns orange when it doesn't add up: Filtered with
no filter reason, Missed with one, or Taken with filter reasons and no Filter notes saying why.
"""
import sys
from pathlib import Path

DECISIONS = ["Filtered", "Missed", "Taken"]
GRADES = ["A+", "A", "B", "C", "D", "F"]
FILTERS = ["Filter 1", "Filter 2", "Filter 3"]
CONFS = ["Confluence 1", "Confluence 2", "Confluence 3"]
REASONS = ["Grade reason 1", "Grade reason 2", "Grade reason 3"]
COLS = (["Decision"] + FILTERS + ["Filter notes"] + CONFS + ["Auto confluence", "Grade", "Rule grade"] + REASONS
        + ["Grade notes", "General notes"])
COMPUTED = ("Rule grade",)              # rewritten on every save
KEEP = [c for c in COLS if c not in COMPUTED]
SHEETS = ("Lists", "Auto confluence")  # kept through extract.py's rebuilds
SEED = {  # generic starters; a setup's own lists come from --from or the user
    "Filter reasons": ["Outside trading hours", "News due", "Against higher-timeframe trend", "No clean setup",
                       "Max trades for the day"],
    "Grade reasons (+ / -)": ["+ With higher-timeframe trend", "+ Clean setup", "+ Room to target", "- Late entry",
                              "- Choppy price action", "- Level in the way"],
    "Confluence": ["VWAP", "50 SMA", "100 SMA", "200 SMA", "SR", "Fib 61.8", "Previous day high", "Previous day low"],
}
AUTO_SEED = [("VWAP", "Yes", "VWAP"), ("SMA 50", "Yes", "50 SMA"), ("SMA 100", "Yes", "100 SMA"),
             ("SMA 200", "Yes", "200 SMA"), ("EMA 21", "No", "21 EMA"), ("Previous day high", "No", "Previous day high"),
             ("Previous day low", "No", "Previous day low")]
AUTO_HEAD_ROW = 5
GUIDE = [
    "How to use",
    "Add a row to any list and it appears in that column's dropdowns (keep each list without gaps).",
    "Grade reasons must start with + (better trade) or - (weaker trade).",
    "Decision: every trade starts as Taken. Change it to Filtered (a rule said no) or Missed (should have taken it). Taken with filter reasons = you took it because something outweighed them: say what in Filter notes.",
    "Confluence = a level close to the entry, at or behind it (a level between entry and target is in the way, not a confluence).",
    "The Auto confluence sheet sets how close counts and which levels are filled in for you; delete one from a trade and it stays deleted.",
    "Grade before you know the result.",
    "",
    "Rule grade (worked out from your entries)",
    "C  = two or more - reasons",
    "B  = one - reason",
    "A+ = no - reasons and 2+ pluses (+ reasons and confluences)",
    "A  = no - reasons",
    "",
    "Safeguards",
    "Decision turns orange when it doesn't add up: Filtered with no filter reason, Missed with one, or Taken with filter reasons and no Filter notes.",
    "Grade turns yellow when it differs from Rule grade. F is never worked out for you: it's your call.",
]
AUTO_NOTES = [
    "Distance: how close to the entry a level must be, at or behind it (on the stop side), to count as a confluence.",
    "Unit: pts = price units (DAX 10 = 10 points); pips = 10 ticks (EURUSD 10 = 0.0010). Leave Distance blank to switch auto-fill off.",
    "Levels: VWAP, SMA <length>, EMA <length> (on the stored 5m bars), Previous day high, Previous day low. Add rows for other lengths.",
    "Indicator <plot name>: any column stored from your exports (e.g. Indicator PDH), its value at the last candle before the entry.",
    "Prior VAH / Prior VAL / Prior POC: the previous session's value area; needs the symbol's value_area setting (config.py symbol SYM value_area).",
    "Name: what goes in the Confluence columns; match it to your Confluence list so the dropdown and the Rule grade agree.",
]


def enabled(path):
    from openpyxl import load_workbook
    return Path(path).exists() and "Lists" in load_workbook(path, read_only=True).sheetnames


def read_lists(path, wb=None):
    """{sheet: (cells, widths)} for the journal's own sheets, to put back after extract.py rebuilds the workbook."""
    from openpyxl import load_workbook
    wb = wb or load_workbook(path)
    out = {}
    for name in SHEETS:
        if name not in wb.sheetnames: continue
        ws = wb[name]
        out[name] = ([(c.row, c.column, c.value) for row in ws.iter_rows() for c in row if c.value is not None],
                     {k: d.width for k, d in ws.column_dimensions.items() if d.width})
    return out


def write_lists(wb, saved=None, distance=None, unit="pts"):
    """Lists and Auto confluence sheets: from `saved` (read_lists) when given, else the generic seeds. The guide in
    Lists column E is always rewritten; an older workbook without an Auto confluence sheet gets one that keeps its
    old behaviour (VWAP and the 50/100/200 SMA within 10 pts)."""
    from openpyxl.styles import Font, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    saved = saved or {}
    ws = wb.create_sheet("Lists")
    if "Lists" in saved:
        cells, widths = saved["Lists"]
        for r, c, v in cells:
            if c != 5: ws.cell(r, c, v)
        for k, w in widths.items(): ws.column_dimensions[k].width = w
    else:
        for j, (head, items) in enumerate(SEED.items(), 1):
            ws.cell(1, j, head)
            for i, v in enumerate(items, 2): ws.cell(i, j, v)
    for k, w in (("A", 34), ("B", 30), ("C", 18), ("D", 3), ("E", 120)): ws.column_dimensions[k].width = w
    for i, v in enumerate(GUIDE, 1):
        ws.cell(i, 5, v)
        if v in ("How to use", "Rule grade (worked out from your entries)", "Safeguards"): ws.cell(i, 5).font = Font(bold=True)
    for j in range(1, 4): ws.cell(1, j).font = Font(bold=True)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for c in row: c.alignment = Alignment(vertical="top")

    au = wb.create_sheet("Auto confluence")
    if "Auto confluence" in saved:
        cells, widths = saved["Auto confluence"]
        for r, c, v in cells:
            if c != 5: au.cell(r, c, v)
    else:
        legacy = "Lists" in saved  # journal made before this sheet existed: keep what it did
        au["A1"], au["B1"] = "Distance", 10 if legacy else distance
        au["A2"], au["B2"] = "Unit", "pts" if legacy else unit
        for j, h in enumerate(("Level", "On", "Name"), 1): au.cell(AUTO_HEAD_ROW, j, h)
        rows = AUTO_SEED if not legacy else [r if r[0] in ("VWAP", "SMA 50", "SMA 100", "SMA 200") else (r[0], "No", r[2]) for r in AUTO_SEED]
        for i, r in enumerate(rows, AUTO_HEAD_ROW + 1):
            for j, v in enumerate(r, 1): au.cell(i, j, v)
    for i, v in enumerate(AUTO_NOTES, 1): au.cell(i, 5, v)
    for c in ("A1", "A2") + tuple(f"{x}{AUTO_HEAD_ROW}" for x in "ABC"): au[c].font = Font(bold=True)
    for k, w in (("A", 20), ("B", 8), ("C", 20), ("D", 3), ("E", 120)): au.column_dimensions[k].width = w
    for formula, rng in (('"pts,pips"', "B2"), ('"Yes,No"', f"B{AUTO_HEAD_ROW + 1}:B{AUTO_HEAD_ROW + 50}")):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True); au.add_data_validation(dv); dv.add(rng)


def auto_settings(wb, symbol):
    """(distance in price units, unit label, unit size, [(kind, length, name)]) from the Auto confluence sheet, or None
    when auto-fill is off (no sheet, no distance, or pips without a known tick)."""
    import re
    sys.path.insert(0, str(Path(__file__).resolve().parent)); import config
    if "Auto confluence" not in wb.sheetnames: return None
    au = wb["Auto confluence"]
    dist, unit = au["B1"].value, str(au["B2"].value or "pts").strip().lower()
    try: dist = float(dist)
    except (TypeError, ValueError): return None
    size = 1.0
    if unit == "pips":
        tick = config.symbol(symbol)["tick"] if symbol else None
        if not tick: return None
        size = tick * 10
    levels = []
    for r in au.iter_rows(min_row=AUTO_HEAD_ROW + 1, max_col=3, values_only=True):
        lv, on, name = r
        if not lv or str(on or "").strip().lower() not in ("yes", "y", "on", "true"): continue
        m = re.fullmatch(r"\s*(SMA|EMA)\s*(\d+)\s*", str(lv), re.I)
        mi = re.fullmatch(r"\s*indicator\s+(.+?)\s*", str(lv), re.I)
        mv = re.fullmatch(r"\s*prior\s+(VAH|VAL|POC)\s*", str(lv), re.I)
        if m: kind = (m.group(1).upper(), int(m.group(2)))
        elif mi: kind = ("indicator", mi.group(1))           # a stored export column, by its plot name
        elif mv: kind = ("prior va", mv.group(1).upper())    # the previous session's value area
        else: kind = (str(lv).strip().lower(), None)
        levels.append((kind[0], kind[1], str(name or lv).strip()))
    return dist * size, unit, size, levels


def autofill_decision(r):
    """Every trade is Taken unless the user (or --decision) says Filtered or Missed."""
    if not r.get("Decision"): r["Decision"] = "Taken"


def style(ws, cols, n_rows):
    """Dropdowns, the Rule grade formula and the live safeguard highlights on the Trades sheet."""
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    L = {c: get_column_letter(cols.index(c) + 1) for c in COLS}
    last = max(n_rows + 1, 2)
    rng = lambda c: f"{L[c]}2:{L[c]}{last}"
    lst = lambda col: f"OFFSET(Lists!${col}$2,0,0,MAX(1,COUNTA(Lists!${col}$2:${col}$1000)),1)"
    for choices, targets, title in ((f'"{",".join(DECISIONS)}"', ["Decision"], "Taken unless you filtered or missed it"),
                                    (lst("A"), FILTERS, "Why you didn't take it"),
                                    (lst("C"), CONFS, "Level at the entry"),
                                    (f'"{",".join(GRADES)}"', ["Grade"], "Grade it before you know the result"),
                                    (lst("B"), REASONS, "+ better / - weaker")):
        dv = DataValidation(type="list", formula1=choices, allow_blank=True, showErrorMessage=False,
                            promptTitle=title, prompt="Pick from the list; add choices on the Lists sheet.")
        ws.add_data_validation(dv)
        for c in targets: dv.add(rng(c))
    f = lambda cs: f"{L[cs[0]]}{{r}}:{L[cs[-1]]}{{r}}"
    filt, conf, rea = f(FILTERS), f(CONFS), f(REASONS)
    for r in range(2, n_rows + 2):
        F_, C_, R_ = filt.format(r=r), conf.format(r=r), rea.format(r=r)
        minus, plus = f'COUNTIF({R_},"-*")', f'(COUNTIF({R_},"+*")+COUNTA({C_}))'
        ws[f"{L['Rule grade']}{r}"] = (
            f'=IF(COUNTA({F_},{C_},{R_},{L["Grade"]}{r})=0,"",'
            f'IF({minus}>=2,"C",IF({minus}=1,"B",IF({plus}>=2,"A+","A"))))')
        ws[f"{L['Rule grade']}{r}"].font = Font(color="555555")
    orange, yellow = PatternFill("solid", bgColor="FFCC80"), PatternFill("solid", bgColor="FFF59D")
    D, F1, F3, FN, G, RG = L["Decision"], L["Filter 1"], L["Filter 3"], L["Filter notes"], L["Grade"], L["Rule grade"]
    n = f"COUNTA(${F1}2:${F3}2)"
    ws.conditional_formatting.add(rng("Decision"), FormulaRule(formula=[
        f'OR(AND(${D}2="Filtered",{n}=0),AND(${D}2="Missed",{n}>0),AND(OR(${D}2="Taken",${D}2=""),{n}>0,${FN}2=""))'], fill=orange))
    ws.conditional_formatting.add(rng("Grade"), FormulaRule(formula=[f'AND(${G}2<>"",${RG}2<>"",${G}2<>${RG}2)'], fill=yellow))
    widths = {"Decision": 10, "Filter notes": 30, "Auto confluence": 22, "Grade": 7, "Rule grade": 7,
              "Grade notes": 30, "General notes": 40, **{c: 22 for c in FILTERS + REASONS}, **{c: 12 for c in CONFS}}
    for c, w in widths.items(): ws.column_dimensions[L[c]].width = w
    for c in COLS: ws.column_dimensions[L[c]].outline_level = 1  # one group: collapse it with the - above the sheet


def auto_confluence(v, levels, max_dist, unit="pts", size=1.0):
    """levels: {name: price}. Returns (Auto confluence text, updated Confluence 1-3) for one trade row `v`.
    A level counts when it is within max_dist (price units) of the entry, at or behind it (the stop side). A level
    found before (listed in the old Auto confluence) but no longer in the slots was removed by the user and is not
    re-added."""
    e, sign = v.get("Entry"), 1 if v.get("Direction") == "Long" else -1
    if e is None: return v.get("Auto confluence"), [v.get(c) for c in CONFS]
    found = [(n, (e - p) * sign) for n, p in levels.items() if p is not None and 0 <= (e - p) * sign <= max_dist + 1e-9]
    before = {s.split(" (")[0].strip() for s in str(v.get("Auto confluence") or "").split(";") if s.strip()}
    slots = [v.get(c) for c in CONFS]
    for n, _ in found:
        if n in slots or n in before or None not in slots: continue
        slots[slots.index(None)] = n
    text = "; ".join(f"{n} ({d / size:.1f}{'' if unit == 'pts' else ' pips'})" for n, d in found) or None
    return text, slots


def init(path, source=None, distance=None, unit="pts"):
    from openpyxl import load_workbook
    path = Path(path)
    wb = load_workbook(path)
    if "Lists" in wb.sheetnames: print(f"already has a Lists sheet: {path}"); return
    saved = read_lists(source) if source else None
    if source and "Lists" not in saved: sys.exit(f"{source} has no Lists sheet to copy")
    write_lists(wb, saved, distance, unit); wb.save(path)
    print(f"Lists and Auto confluence sheets added: {path}" + (f" (copied from {Path(source).name})" if source else ""))
    if not source and distance is None: print("Auto confluence is off until a Distance is set on its sheet (or pass --distance).")


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(usage=__doc__); ap.add_argument("cmd", choices=["init"]); ap.add_argument("workbook")
    ap.add_argument("--from", dest="source"); ap.add_argument("--distance", type=float)
    ap.add_argument("--unit", choices=["pts", "pips"], default="pts")
    a = ap.parse_args()
    init(a.workbook, a.source, a.distance, a.unit)
