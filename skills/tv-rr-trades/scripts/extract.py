"""Extract TradingView Long/Short Position (risk/reward) drawings from the clipboard into an Excel trade log.

usage: extract.py [--file CLIP.html] [--tick 0.1] [--out PATH] [--dry-run] [--recheck]
                  [--decision Filtered|Missed|Taken] [--remove]

--decision sets the journal's Decision on every drawing on the clipboard: drawings not in the sheet are added, ones
already there are updated ("add these filtered trades", "update these trades to missed"). Taken clears it (blank =
taken). The workbook needs journal columns (journal.py init).
--remove deletes the clipboard's drawings from the sheet ("remove these trades"); with --dry-run it only lists them.

TradingView copies drawings as HTML: one <span data-tradingview-clip="…"> whose JSON holds a `sources` list, one per
drawing. For LineToolRiskRewardLong / LineToolRiskRewardShort:
  points[0]          entry time and price
  points[1]          right edge of the box — NOT an exit
  points[3]          when present: where TP or stop was hit (exit time and price)
  state.stopLevel / state.profitLevel   stop and target distance in TICKS
The tick size is confirmed from the exit point where one exists, else read from the config or --tick; it is never
guessed from the entry's decimals (25172 does not mean a tick of 1). Settings: config.py. Prices: prices.py.
"""
import argparse, datetime as dt, html, json, math, platform, re, subprocess, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config, journal, prices  # noqa: E402
from prices import bars_for  # noqa: E402,F401  (re-exported for the analysis skill)

TZ, UTC = config.tz(), dt.timezone.utc
ANALYSIS = Path(__file__).resolve().parents[2] / "tv-rr-analysis" / "scripts"
RR_TYPES = {"LineToolRiskRewardLong": "Long", "LineToolRiskRewardShort": "Short"}
COLS = ["Symbol", "Direction", "Entry time", "Chart", "TV chart", "Outcome", "Timeframe", "Entry", "Stop", "TP planned", "Risk pts",
        "Reward pts", "Planned R:R", "Status", "Fill time", "Exit time", "Exit price", "Result pts", "Result R",
        "Check", "Confirmed outcome", "Note",
        "TradingView says", "TradingView exit", "Checked on", "Checked to", "Drawing id", "Entry (UTC)", "Last copied"]
# Local-time columns, shown in the configured timezone. Older workbooks named them "Entry time (Sydney)" etc.
TIME_COLS = ("Entry time", "Fill time", "Exit time", "TradingView exit", "Checked to")
# Filled in by the user in Excel; the script never overwrites them
USER_COLS = ("Confirmed outcome",)
# Filled in by the user and kept whatever happens to the drawing (a TradingView chart link: a URL, or a link with text)
KEEP_COLS = ("TV chart",)
CONFIRM_CHOICES = ["TP", "Stop", "Not filled", "Open"]


def local(t):
    return dt.datetime.fromtimestamp(t, TZ).replace(tzinfo=None)


def read_clipboard():
    """The clipboard's HTML flavour, where TradingView puts the drawings (its plain text is just "Drawings")."""
    system = platform.system()
    try:
        if system == "Darwin":
            r = subprocess.run(["osascript", "-e", "the clipboard as «class HTML»"], capture_output=True, text=True)
            m = re.search(r"«data HTML([0-9A-Fa-f]*)»", r.stdout)
            return bytes.fromhex(m.group(1)).decode("utf-8", "replace") if m else ""
        if system == "Windows":
            ps = "Add-Type -AssemblyName System.Windows.Forms; [System.Windows.Forms.Clipboard]::GetText('Html')"
            return subprocess.run(["powershell", "-NoProfile", "-STA", "-Command", ps], capture_output=True, text=True,
                                  encoding="utf-8").stdout
        for cmd in (["wl-paste", "-t", "text/html"], ["xclip", "-selection", "clipboard", "-t", "text/html", "-o"]):
            try: return subprocess.run(cmd, capture_output=True, text=True).stdout
            except FileNotFoundError: continue
    except FileNotFoundError:
        pass
    return ""


def drawings(text):
    out, skipped = [], 0
    for raw in re.findall(r'data-tradingview-clip="([^"]*)"', text):
        for src in json.loads(html.unescape(raw)).get("sources", []):
            so = src.get("source", {})
            if so.get("type") in RR_TYPES: out.append(so)
            else: skipped += 1
    return out, skipped


def pow10(x):
    if not x or x <= 0: return None
    e = round(math.log10(x))
    return 10.0 ** e if abs(x - 10.0 ** e) <= 1e-6 * 10.0 ** e else None


def decimals(tick):
    return max(0, -int(round(math.log10(tick))))


def build_row(so, tick_override, now):
    st, pts = so.get("state", {}), so.get("points", [])
    side = RR_TYPES[so["type"]]; sign = 1 if side == "Long" else -1
    sym = st.get("symbol", ""); entry = float(pts[0]["price"]); t0 = int(pts[0]["time_t"])
    sl_ticks, tp_ticks = float(st.get("stopLevel") or 0), float(st.get("profitLevel") or 0)
    close = pts[3] if len(pts) >= 4 else None
    tick, learned = tick_override, None
    if tick is None and close:
        d = abs(float(close["price"]) - entry)
        for lv in (tp_ticks, sl_ticks):
            cand = pow10(d / lv) if lv else None
            if cand: tick = learned = cand; break
    if tick is None: tick = config.symbol(sym)["tick"]
    row = dict(zip(COLS, [None] * len(COLS)))
    row.update({"Symbol": sym, "Timeframe": st.get("interval"), "Direction": side, "Entry": entry, "Drawing id": so.get("id"),
                "Entry time": local(t0),
                "Entry (UTC)": dt.datetime.fromtimestamp(t0, UTC).replace(tzinfo=None), "Last copied": now})
    warn = None
    if tick:
        nd = decimals(tick)
        stop, tp = round(entry - sign * sl_ticks * tick, nd), round(entry + sign * tp_ticks * tick, nd)
        risk, reward = round(sl_ticks * tick, nd), round(tp_ticks * tick, nd)
        row.update({"Stop": stop, "TP planned": tp, "Risk pts": risk, "Reward pts": reward,
                    "Planned R:R": round(reward / risk, 2) if risk else None})
    else:
        warn = f"{sym}: tick size unknown — stop/TP left blank (pass --tick, or: config.py symbol {sym} tick 0.1)"
    if close:
        px = float(close["price"])
        row.update({"Exit time": local(int(close["time_t"])), "Exit price": px})
        if tick and abs(px - row["TP planned"]) <= tick / 2: out = "TP"
        elif tick and abs(px - row["Stop"]) <= tick / 2: out = "Stop"
        else: out = "Closed (other)"
        row["Outcome"] = row["TradingView says"] = out
        row["TradingView exit"] = row["Exit time"]
        res = round((px - entry) * sign, decimals(tick) if tick else 2)
        row["Result pts"] = res
        row["Result R"] = round(res / row["Risk pts"], 2) if row.get("Risk pts") else None
    else:
        row["Outcome"] = row["TradingView says"] = "Not closed"
    row["Status"] = "Closed" if close else "Not closed"
    return row, learned, warn


def candle_path(b):
    """Likely intrabar path of a candle: bullish O→L→H→C, bearish O→H→L→C."""
    return [b["o"], b["l"], b["h"], b["c"]] if b["c"] >= b["o"] else [b["o"], b["h"], b["l"], b["c"]]


def simulate(bars, t0, long_, e, sl, tp):
    """Walk each candle's likely path from the entry candle on: fill at the entry price, then whichever of stop or TP
    the path reaches first. Returns (status, outcome, fill_bar, exit_bar, exit_price, inferred) — inferred is True when
    the call depended on the order inside one candle (fill and exit in the same candle, or both levels in one)."""
    filled, fill_b = False, None
    for b in (x for x in bars if x["t"] >= t0):
        p = candle_path(b); start = 0
        if not filled:
            for k in range(3):
                lo, hi = sorted((p[k], p[k + 1]))
                if lo <= e <= hi: filled, fill_b, start = True, b, k; break
            if not filled: continue
        cur = e if b is fill_b else p[0]
        legs = [(cur, p[start + 1])] + [(p[k], p[k + 1]) for k in range(start + 1, 3)] if b is fill_b else \
               [(p[k], p[k + 1]) for k in range(3)]
        both = ((b["l"] <= sl <= b["h"]) and (b["l"] <= tp <= b["h"]))
        for a, z in legs:
            lo, hi = sorted((a, z)); hits = [(abs(lv - a), name, lv) for name, lv in (("Stop", sl), ("TP", tp)) if lo <= lv <= hi]
            if hits:
                _, name, lv = min(hits)
                return "Closed", name, fill_b, b, lv, (b is fill_b and not certain(bars, fill_b, sl, tp, name)) or both
    return ("Open" if filled else "Not filled"), ("Open" if filled else "Not filled"), fill_b, None, None, False


def certain(bars, fill_b, sl, tp, outcome):
    """An exit decided in the fill candle is still certain when the order inside that candle cannot change it:
    if the candle reached only one of the levels, the trade either hit it after the fill or carried on to the next
    candles — so when those next candles reach the same level first, the outcome is the same either way
    (28 Jul 2026: entry and TP in one candle, no stop; the next candle reached the TP)."""
    in_fill = [n for n, lv in (("Stop", sl), ("TP", tp)) if fill_b["l"] <= lv <= fill_b["h"]]
    if len(in_fill) != 1 or in_fill[0] != outcome: return False
    for b in (x for x in bars if x["t"] > fill_b["t"]):
        s_, t_ = b["l"] <= sl <= b["h"], b["l"] <= tp <= b["h"]
        if s_ and t_: return False
        if s_ or t_: return ("Stop" if s_ else "TP") == outcome
    return False


def check_with_prices(r):
    """TradingView's close point is the record, except when it falls in the entry candle: TradingView counts a level
    anywhere in that candle's range, even prices that traded before the fill (29 Sep 2026: TP 'hit' at 17:20 was the
    candle's open, before the short filled; the trade was stopped at 17:25). Those, and drawings TradingView shows as
    not closed, are re-derived from prices along each candle's likely path (1m wherever it covers the trade)."""
    tv = r.get("TradingView says")
    t0 = int(r["Entry (UTC)"].replace(tzinfo=UTC).timestamp()) if r.get("Entry (UTC)") else None
    bars, tf = bars_for(r["Symbol"], t0)
    r["Checked on"] = tf
    if not bars or r.get("Stop") is None or r.get("TP planned") is None:
        if tv not in (None, "Not closed"): r.update({"Status": "Closed", "Outcome": tv, "Note": "exit from TradingView"})
        return r
    long_ = r["Direction"] == "Long"; e, sl, tp = r["Entry"], r["Stop"], r["TP planned"]
    last = local(bars[-1]["t"])
    syd = lambda b: local(b["t"])
    r["Checked to"] = last
    tv_exit = r.get("TradingView exit")
    entry_candle = local(t0)
    if tf == "1m":
        in_fill = False
    elif tv in ("TP", "Stop") and tv_exit and tv_exit > entry_candle:
        fill_b = simulate(bars, t0, long_, e, sl, tp)[2]
        if fill_b is None or tv_exit > syd(fill_b):
            r.update({"Status": "Closed", "Outcome": tv, "Exit time": tv_exit, "Note": "exit from TradingView"})
            return r
        in_fill = True
    else:
        in_fill = False
    if bars[-1]["t"] < t0:
        r.update({"Status": "No price data yet", "Outcome": tv or "Not closed", "Note": "price data ends before the entry — add newer exports (prices.py ingest) and rerun"})
        return r
    status, out, fb, xb, px, inferred = simulate(bars, t0, long_, e, sl, tp)
    for k in ("Fill time", "Exit time", "Exit price", "Result pts", "Result R"): r[k] = None
    notes = []
    if tf == "1m":
        notes.append("checked on 1m prices")
        if tv in ("TP", "Stop") and status == "Closed" and out != tv: notes.append(f"TradingView shows {tv}")
        elif tv == "Closed (other)": notes.append("TradingView's close is at neither the TP nor the stop")
    elif tv == "Closed (other)": notes.append(f"TradingView's close is at neither the TP nor the stop; re-checked on {tf} prices")
    elif tv not in (None, "Not closed"):
        notes.append(f"TradingView shows {tv} in the {'fill' if in_fill else 'entry'} candle; re-checked on {tf} prices")
    else: notes.append(f"TradingView shows no exit; checked on {tf} prices")
    r["Fill time"] = syd(fb) if fb else None
    if status == "Closed":
        res = round((px - e) * (1 if long_ else -1), config.decimals(r["Symbol"], e))
        r.update({"Status": "Closed", "Outcome": out, "Exit time": syd(xb), "Exit price": px, "Result pts": res,
                  "Result R": round(res / r["Risk pts"], 2) if r.get("Risk pts") else None})
        if inferred: notes.append(f"order inside one {tf} candle inferred from its open/close" + (" — check on a 1m chart" if tf != "1m" else ""))
    elif status == "Not filled":
        r.update({"Status": "Not filled", "Outcome": "Not filled"}); notes.append(f"price never traded at the entry (checked to {last:%a %d %b %H:%M})")
    else:
        r.update({"Status": "Open", "Outcome": "Open"}); notes.append(f"no TP or stop hit up to {last:%a %d %b %H:%M}")
    r["Note"] = "; ".join(notes)
    return r


def apply_confirmation(r):
    """Flag results that still need the user's eye. Checked on 1m prices: only a call decided inside one 1m candle, or
    a disagreement with TradingView. Checked on coarser bars (no 1m coverage): anything the script worked out itself. A confirmed
    outcome overrides the computed one, and its result is re-scored."""
    note = r.get("Note") or ""
    tv, out = r.get("TradingView says"), r.get("Outcome")
    reasons = []
    if "checked on 1m prices" in note:
        if "order inside one 1m candle" in note: reasons.append("decided inside one 1m candle")
    elif "neither the TP nor the stop" in note: reasons.append("TradingView's close is at neither the TP nor the stop")
    elif "in the fill candle; re-checked" in note: reasons.append("TradingView's exit was in the candle the order filled in")
    elif "re-checked on 5m prices" in note: reasons.append("TradingView's exit was in the entry candle")
    elif "TradingView shows no exit" in note: reasons.append("TradingView shows no exit")
    m = re.search(r"order inside one (\d+m) candle", note)
    if m and m.group(1) != "1m": reasons.append(f"decided inside one {m.group(1)} candle")
    if tv in ("TP", "Stop") and out in ("TP", "Stop", "Both in one candle") and out != tv: reasons.append(f"TradingView says {tv}")
    conf = r.get("Confirmed outcome")
    if conf:
        if conf != r.get("Outcome"): r["Exit time"] = None  # the computed exit time was for a different outcome
        r["Outcome"] = conf
        risk, reward = r.get("Risk pts"), r.get("Reward pts")
        if conf == "TP" and reward is not None: r.update({"Result pts": reward, "Result R": round(reward / risk, 2) if risk else None, "Exit price": r.get("TP planned")})
        elif conf == "Stop" and risk is not None: r.update({"Result pts": -risk, "Result R": -1.0, "Exit price": r.get("Stop")})
        else: r.update({"Result pts": None, "Result R": None})
        if conf in ("Not filled", "Open"): r.update({"Exit price": None, "Exit time": None})
        r["Check"] = "✓ confirmed"
    else:
        r["Check"] = "CONFIRM: " + "; ".join(reasons) if reasons else None
    return r


def key(r):
    # the drawing id survives moving the drawing (new entry time, other timeframe), so it alone identifies the trade
    return (str(r["Symbol"]), str(r["Drawing id"]))


def load(path):
    from openpyxl import load_workbook
    if not path.exists(): return []
    ws = load_workbook(path)["Trades"]; rows = list(ws.iter_rows(values_only=True))
    head = [legacy_name(h) for h in rows[0]]
    out = []
    for n, r in enumerate(rows[1:], 2):
        d = dict(zip(head, r))
        if not d.get("Symbol"): continue  # a row is a trade only with a symbol
        for c in KEEP_COLS:  # a link's target, when the cell shows text over it
            link = ws.cell(n, head.index(c) + 1).hyperlink if c in head else None
            if link is not None and link.target: d[c] = link.target
        out.append(d)
    # times copied from TradingView are stored as local times: move them if the timezone setting has changed
    note = ws.cell(1, head.index("Entry time") + 1).comment if "Entry time" in head else None
    m = re.search(r"shown in (\S+)\.", note.text) if note else None
    old = m.group(1) if m else ("Australia/Sydney" if any("(Sydney)" in str(h) for h in rows[0]) else None)
    if old and old != config.tz_name():
        from zoneinfo import ZoneInfo
        for r in out:
            if isinstance(r.get("TradingView exit"), dt.datetime):
                r["TradingView exit"] = r["TradingView exit"].replace(tzinfo=ZoneInfo(old)).astimezone(TZ).replace(tzinfo=None)
    return out


def legacy_name(h):
    """'Entry time (Sydney)' → 'Entry time': older workbooks put the timezone in the column name."""
    m = re.fullmatch(r"(.+) \((?!UTC\))[^)]+\)", str(h or ""))
    return m.group(1) if m and m.group(1) in TIME_COLS else h


def columns(with_journal):
    """COLS, with the journal columns after them (after "Last copied") when the workbook has a Lists sheet (journal.py)."""
    return list(COLS) + (journal.COLS if with_journal else [])


def save(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    lists = journal.read_lists(path) if journal.enabled(path) else None
    COLS = columns(lists is not None)
    wb = Workbook(); ws = wb.active; ws.title = "Trades"
    ws.append(COLS)
    for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "G2"  # header row and Symbol, Direction, Entry time, Chart, TV chart, Outcome stay in view
    from openpyxl.comments import Comment
    ws.cell(1, COLS.index("Entry time") + 1).comment = Comment(f"Local times are shown in {config.tz_name()}.", "tv-rr")
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"' + ",".join(CONFIRM_CHOICES) + '"', allow_blank=True,
                        promptTitle="Confirm the outcome", prompt="Pick what really happened; leave blank to keep the computed result.")
    ws.add_data_validation(dv)
    fills = {"TP": "C8E6C9", "Stop": "FFCDD2", "Not closed": "E0E0E0", "Not filled": "E0E0E0", "Open": "BBDEFB",
             "Closed (other)": "FFF9C4", "Both in one candle": "FFF9C4"}
    for r in rows:
        r["Chart"] = None  # links to the gallery, re-made by tv-rr-analysis after every save
        ws.append([r.get(c) for c in COLS])
        i = ws.max_row
        for c in TIME_COLS + ("Entry (UTC)", "Last copied"):
            ws.cell(i, COLS.index(c) + 1).number_format = "ddd dd mmm yyyy hh:mm"
        for c in ("Entry", "Stop", "TP planned", "Exit price"):
            ws.cell(i, COLS.index(c) + 1).number_format = "#,##0.0#####"
        tc = ws.cell(i, COLS.index("TV chart") + 1)
        if str(r.get("TV chart") or "").startswith(("http://", "https://")):
            tc.value = "TV chart"; tc.hyperlink = r["TV chart"]; tc.font = Font(color="0563C1", underline="single")
        o = ws.cell(i, COLS.index("Outcome") + 1)
        if r.get("Outcome") in fills: o.fill = PatternFill("solid", fgColor=fills[r["Outcome"]])
        cc = ws.cell(i, COLS.index("Confirmed outcome") + 1); dv.add(cc)
        ck = ws.cell(i, COLS.index("Check") + 1)
        if str(r.get("Check") or "").startswith("CONFIRM"):
            ck.fill = cc.fill = PatternFill("solid", fgColor="FFE082"); ck.font = Font(bold=True)
        elif r.get("Check"): ck.font = Font(color="2E7D32")
    widths = {"Symbol": 16, "Timeframe": 10, "Direction": 10, **{c: 22 for c in TIME_COLS}, "Entry (UTC)": 22, "Last copied": 22,
              "Outcome": 16, "Status": 16, "TradingView says": 14,
              "Note": 60, "Drawing id": 12, "Chart": 11, "TV chart": 11, "Check": 44, "Confirmed outcome": 16}
    for i, c in enumerate(COLS, 1): ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 12)
    if lists is not None:
        journal.style(ws, COLS, len(rows)); journal.write_lists(wb, lists)
    ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)


def fmt(r):
    t = r["Entry time"]; x = r["Exit time"]
    if r.get("Status") in ("Open", "Not filled", "No price data yet"):
        sl = f"{r['Stop']:,}" if r["Stop"] is not None else "?"; tp = f"{r['TP planned']:,}" if r["TP planned"] is not None else "?"
        return f"{r['Direction']:5} {r['Symbol']} {t:%a %d %b %Y %H:%M}  entry {r['Entry']:,}  stop {sl}  TP {tp} → {r['Status']}" + (f"  [{r['Note']}]" if r.get("Note") else "")
    sl = f"{r['Stop']:,}" if r["Stop"] is not None else "?"; tp = f"{r['TP planned']:,}" if r["TP planned"] is not None else "?"
    ex = f" → exit {x:%a %d %b %H:%M} at {r['Exit price']:,} ({r['Outcome']}, {r['Result R']:+}R)" if x and r["Result R"] is not None else \
        f" → exit {x:%a %d %b %H:%M} ({r['Outcome']})" if x else " → not closed"
    ex += f"  [{r['Note']}]" if r.get("Note") else ""
    return f"{r['Direction']:5} {r['Symbol']} {t:%a %d %b %Y %H:%M}  entry {r['Entry']:,}  stop {sl}  TP {tp}{ex}"


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--file"); ap.add_argument("--tick", type=float)
    ap.add_argument("--out"); ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--recheck", action="store_true", help="re-assess the trades already in the workbook; no clipboard")
    ap.add_argument("--decision", choices=["Filtered", "Missed", "Taken"], help="set Decision on the copied drawings (Taken clears it)")
    ap.add_argument("--remove", action="store_true", help="delete the copied drawings from the sheet")
    a = ap.parse_args()
    path = config.workbook(a.out)
    if a.decision and not journal.enabled(path):
        sys.exit(f"{path.name} has no journal columns, so there is no Decision to set. Add them first: "
                 f"journal.py init \"{path}\", then run again.")
    if config.load()["exports_dirs"]: prices.ingest(verbose=False)
    text = "" if a.recheck else Path(a.file).read_text() if a.file else read_clipboard()
    found, skipped = drawings(text)
    if not found and not a.recheck:
        sys.exit("No TradingView risk/reward drawings on the clipboard. In TradingView, select the Long/Short Position "
                 "tools, copy them (Cmd-C / Ctrl-C), then run again." + (f" ({skipped} other drawings were skipped.)" if skipped else ""))
    now = dt.datetime.now(TZ).replace(tzinfo=None, microsecond=0)
    new, warns, learned = [], [], {}
    for so in found:
        r, lt, w = build_row(so, a.tick, now); new.append(r)
        if lt and not config.symbol(r["Symbol"])["tick"]: learned[r["Symbol"]] = lt
        if w: warns.append(w)
    old = {key(r): r for r in load(path)}
    if a.remove:
        gone = [old.pop(key(r)) for r in new if key(r) in old]; missing = [r for r in new if key(r) not in {key(g) for g in gone}]
        print(f"{len(found)} drawings on the clipboard: {len(gone)} in {path.name}"
              + (", nothing to remove" if not gone else " — dry run, nothing removed" if a.dry_run else ", removed"))
        for r in sorted(gone, key=lambda r: r["Entry (UTC)"]): print("  - " + fmt(r))
        for r in missing: print(f"  not in the sheet: {r['Direction']} {r['Symbol']} {r['Entry time']:%a %d %b %Y %H:%M}")
        if gone and not a.dry_run:
            save(path, sorted(old.values(), key=lambda r: r["Entry (UTC)"]))
            if (ANALYSIS / "enrich.py").exists():
                sys.path.insert(0, str(ANALYSIS)); from enrich import enrich
                enrich(path)
            print(f"saved: {path} ({len(old)} trades)")
        sys.exit(0)
    added = sum(1 for r in new if key(r) not in old); updated = len(new) - added
    for r in new:
        prev = old.get(key(r))
        if prev:
            same = all(prev.get(c) == r.get(c) for c in ("Entry", "Stop", "TP planned"))
            for c in USER_COLS: r[c] = prev.get(c) if same else None
            for c in journal.KEEP + list(KEEP_COLS): r[c] = prev.get(c)  # your notes and links stay even when the levels move
            if not same and prev.get("Confirmed outcome"): r["Note"] = "levels changed since you confirmed it — confirm again"
        if a.decision: r["Decision"] = None if a.decision == "Taken" else a.decision
        old[key(r)] = r
    rows = sorted(old.values(), key=lambda r: r["Entry (UTC)"])
    for r in rows:
        r.setdefault("TradingView says", r.get("Outcome")); r.setdefault("TradingView exit", None)
        if r["TradingView exit"] is None and r["TradingView says"] not in (None, "Not closed"):
            r["TradingView exit"] = r.get("Exit time")  # rows written before this column existed
        if r.get("Entry (UTC)"): r["Entry time"] = local(int(r["Entry (UTC)"].replace(tzinfo=UTC).timestamp()))
        check_with_prices(r); apply_confirmation(r); journal.autofill_decision(r)
    print(f"{len(found)} risk/reward drawings on the clipboard ({skipped} other drawings skipped): {added} new, {updated} updated"
          + (" — dry run, nothing written" if a.dry_run else ""))
    print("updated from the clipboard:")
    for r in sorted(new, key=lambda r: r["Entry (UTC)"]): print("  " + fmt(r) + (f"  → {a.decision}" if a.decision else ""))
    todo = [r for r in rows if str(r.get("Check") or "").startswith("CONFIRM")]
    if todo:
        print(f"{len(todo)} to confirm in Excel (set 'Confirmed outcome'):")
        for r in todo: print(f"  {r['Direction']:5} {r['Symbol']} {r['Entry time']:%a %d %b %H:%M}  {r['Outcome']}  — {r['Check'][9:]}")
    stale = [r for r in rows if r.get("Status") in ("Open", "No price data yet")]
    if stale: print("still open / waiting for price data:"); [print("  " + fmt(r)) for r in stale]
    for w in sorted(set(warns)): print("WARNING: " + w)
    if not a.dry_run:
        save(path, rows)
        for sym, t in learned.items(): config.set_symbol(sym, "tick", t)
        print(f"saved: {path} ({len(rows)} trades)")
        if (ANALYSIS / "enrich.py").exists():  # this rebuild drops the analysis columns, so put them straight back
            sys.path.insert(0, str(ANALYSIS)); from enrich import enrich
            enrich(path); print("analysis columns and Breakdown sheet refreshed (tv-rr-analysis)")
