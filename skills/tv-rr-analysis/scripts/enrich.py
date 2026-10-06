"""Add what-if columns and a Breakdown sheet to the TradingView RR trade log, from the stored price data.

usage: enrich.py [--file PATH]

Per trade, on the Trades sheet:
  1/2 stop          same entry and TP, stop at half the planned distance → Win / Loss
  TP 0.5R … 2.5R    same entry and planned stop, TP moved to 0.5/0.75/1/1.5/2/2.5 times the stop distance → Win / Loss
  MAE pts           worst move against the entry between fill and exit; blank when the exit candle reached both the
                    stop and the TP (the order inside it is unknown)
  MFE pts           best move in favour of the entry between fill and exit (for a TP that is the target itself; for a
                    stopped trade, how close it came); blank in the same case
  BE at 55%         the planned trade with the stop moved to the entry once price has gone 55% of the way to the TP:
                    "BE" when price then came back to the entry before the TP (flagged amber), else Win / Loss / Open
  1.5R BE at 55%    the same, with the TP at 1.5R (stop to entry after 0.825R in favour)
  1.5R BE at 1R     TP at 1.5R, stop to entry once price is 1R in favour
  ATR 5m … ATR D    ATR(14) on 5m, 15m, 4h and daily bars at the entry time (last closed bar of each; indicators.py)
  ATR 1:1 … 1.5:1.5 same entry, stop and TP sized from the 5m ATR at entry (stop multiple : TP multiple) → Win / Loss
  VWAP …            the chart's session VWAP at the entry (last closed 5m bar): its value, the entry's distance from it
                    (+ = beyond it your way), which side price closed on, whether it sits between the entry and the TP
                    (in the way) or between the stop and the entry (behind), and whether it is higher or lower than
                    30 minutes earlier (VWAP 30m change — not a slope: it often reaches back before the open's spike)
  VWAP direction    the line's slope at entry, from the last 10 minutes (two closed 5m bars): Flat when it moved less
                    than FLAT_PTS; "Upwards turning" / "Downwards turning" when it still moves that way but at less
                    than half the pace of the 10 minutes before (flattening out); otherwise Upwards / Downwards
  VWAP immediate slope   that 10-minute slope against the trade: With / Against / Flat
  TV <name>         every other indicator column on the exported chart (tv-rr-trades prices.py keeps them), at the last
                    bar closed before the entry, on the trade's bar size when exported, else the smallest held.
                    Grouped at the end of the sheet.
"Open" = neither level reached in the price data yet. Blank = not filled, no price data, or levels unknown.
Breakdown sheet: per variant (Planned, 1/2 stop, 0.5R … 2.5R targets, break-even and ATR versions) and per direction (All / Long / Short) — trades, wins,
losses, open, win %, net points, net R, and MAE for winners and losers.
Uses the same fill rule and candle path as the tv-rr-trades skill (bullish O→L→H→C, bearish O→H→L→C), on 1m bars
where they cover the trade and the finest stored bars otherwise.
"""
import argparse, datetime as dt, statistics, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "tv-rr-trades" / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import config, indicators, journal, prices  # noqa: E402
from extract import UTC, bars_for, candle_path, legacy_name  # noqa: E402

TARGETS = [0.5, 0.75, 1, 1.5, 2, 2.5]
BE_AT = 0.55
ATR_COLS = ["ATR 5m", "ATR 15m", "ATR 4h", "ATR D"]
ATR_RR = [(1, 1), (1, 1.5), (1, 2), (1.5, 1.5)]
ATR_RR_COLS = [f"ATR {a:g}:{b:g}" for a, b in ATR_RR]
VWAP_COLS = ["VWAP", "VWAP dist", "VWAP side", "VWAP in the way", "VWAP behind", "VWAP 30m change", "VWAP direction",
             "VWAP immediate slope"]
RENAMED = {"VWAP slope": "VWAP 30m change"}  # its 30-minute comparison was never the line's slope at entry
FLAT_PTS = 2       # VWAP moving less than this over the last 10 minutes counts as flat
TURNING = 0.5      # still moving the same way, but at less than this share of the previous 10 minutes' pace
TARGET_COLS = [f"TP {t:g}R" for t in TARGETS]
NEW_COLS = ["1/2 stop"] + TARGET_COLS + ["MAE pts", "MFE pts", "BE at 55%", "1.5R BE at 55%", "1.5R BE at 1R"] + ATR_COLS + ATR_RR_COLS + VWAP_COLS
# Earlier versions named the targets in points; their columns are removed so the sheet doesn't carry both
OBSOLETE = ["10pt", "15pt", "20pt", "30pt", "40pt", "50pt", "30pt BE at 55%", "30pt BE at 20pt"]


def walk(bars, t0, long_, e, sl, tp, be_at=None):
    """Fill at the entry along each candle's path, then the first of stop/TP. Returns (result, mae, both_in_exit_candle, mfe):
    result is "Win", "Loss", "Open" or None (never filled); mae / mfe are the worst move against and the best move for the
    entry from the fill to the exit.
    With be_at, the stop moves to the entry once price has gone that fraction of the way to the TP, and a return to
    the entry after that is "BE"."""
    sign = 1 if long_ else -1
    adverse = lambda price: max(0.0, (e - price) * sign)
    favour = lambda price: max(0.0, (price - e) * sign)
    filled, mae, mfe = False, 0.0, 0.0
    trigger = e + sign * be_at * abs(tp - e) if be_at else None; at_be = False
    for b in (x for x in bars if x["t"] >= t0):
        p = candle_path(b)
        if not filled:
            k = next((k for k in range(3) if min(p[k], p[k + 1]) <= e <= max(p[k], p[k + 1])), None)
            if k is None: continue
            filled = True; legs = [(e, p[k + 1])] + [(p[j], p[j + 1]) for j in range(k + 1, 3)]
        else:
            legs = [(p[j], p[j + 1]) for j in range(3)]
        both = b["l"] <= sl <= b["h"] and b["l"] <= tp <= b["h"]
        for a, z in legs:
            lo, hi = min(a, z), max(a, z)
            stop_name, stop_lv = ("BE", e) if at_be else ("Loss", sl)
            hits = [(abs(lv - a), name, lv) for name, lv in ((stop_name, stop_lv), ("Win", tp)) if lo <= lv <= hi and not (name == "BE" and lv == a)]
            if hits:
                _, name, lv = min(hits)
                mae = max(mae, adverse(a), adverse(lv)); mfe = max(mfe, favour(a), favour(lv))
                return name, mae, both, mfe
            mae = max(mae, adverse(a), adverse(z)); mfe = max(mfe, favour(a), favour(z))
            if trigger is not None and not at_be and lo <= trigger <= hi: at_be = True
    return ("Open" if filled else None), (mae if filled else None), False, (mfe if filled else None)


TV_PREFIX = "TV "


def tv_values(ws, head):
    """{sheet row: {"TV <plot name>": value}} from the indicator columns stored with the price exports, read at the last
    bar closed before each entry (no peeking), on the trade's bar size when it was exported."""
    out = {}
    if "Entry (UTC)" not in head: return out
    ce, cs, ct = head.index("Entry (UTC)") + 1, head.index("Symbol") + 1, (head.index("Timeframe") + 1 if "Timeframe" in head else None)
    for i in range(2, ws.max_row + 1):
        e, sym = ws.cell(i, ce).value, ws.cell(i, cs).value
        if not e or not sym: continue
        tf = ws.cell(i, ct).value if ct else None
        try: tf = int(str(tf).rstrip("mM"))
        except (TypeError, ValueError): tf = None
        m, vals = prices.indicators_at(sym, int(e.replace(tzinfo=UTC).timestamp()), tf)
        if vals: out[i] = {TV_PREFIX + tv_name(sym, m, k): v for k, v in vals.items()}
    return out


def tv_name(sym, m, name):
    """The column name for an exported plot: the user's rename (config indicator_names) first; else a plain "EMA" /
    "MA" plot gets its length worked out from the prices ("EMA (2)" → "EMA 20"); else the plot's own name."""
    renames = config.load().get("indicator_names") or {}
    if name in renames: return renames[name]
    found = indicators.ma_length(sym, m, name, prices.indicator_rows(sym, m))
    return f"{found[0]} {found[1]}" if found else name


def vwap_direction(recent, sign):
    """(direction, immediate slope) from [VWAP now, 10 min ago, 20 min ago] and the trade's sign (+1 long, -1 short)."""
    if not recent or recent[0] is None or recent[1] is None: return None, None
    now = recent[0] - recent[1]
    before = recent[1] - recent[2] if recent[2] is not None else None
    if abs(now) < FLAT_PTS: return "Flat", "Flat"
    word = "Upwards" if now > 0 else "Downwards"
    if before is not None and before * now > 0 and abs(now) < TURNING * abs(before): word += " turning"
    return word, "With" if now * sign > 0 else "Against"


def enrich(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path); ws = wb["Trades"]
    for c in ws[1]: c.value = RENAMED.get(legacy_name(c.value), legacy_name(c.value))
    for i in sorted((n for n, c in enumerate(ws[1], 1) if c.value in OBSOLETE), reverse=True): ws.delete_cols(i)
    head = [c.value for c in ws[1]]
    for c in NEW_COLS:
        if c not in head:
            ws.cell(1, len(head) + 1, c).font = Font(bold=True); ws.cell(1, len(head) + 1).alignment = Alignment(horizontal="center")
            head.append(c)
    tv = tv_values(ws, head)  # {row: {"TV <name>": value}}
    now = {c for vals in tv.values() for c in vals}
    for i in sorted((n for n, h in enumerate(head, 1) if str(h).startswith(TV_PREFIX) and h not in now), reverse=True):
        ws.delete_cols(i); del head[i - 1]
    for c in sorted({c for vals in tv.values() for c in vals}, key=str.lower):
        if c not in head:
            ws.cell(1, len(head) + 1, c).font = Font(bold=True); ws.cell(1, len(head) + 1).alignment = Alignment(horizontal="center", wrap_text=True)
            head.append(c)
    col = {h: i + 1 for i, h in enumerate(head)}
    for i in range(2, ws.max_row + 1):
        for h in head:
            if str(h).startswith(TV_PREFIX): ws.cell(i, col[h], tv.get(i, {}).get(h))
    for h in head:
        if str(h).startswith(TV_PREFIX):
            ws.column_dimensions[get_column_letter(col[h])].width = 10; ws.column_dimensions[get_column_letter(col[h])].outline_level = 1
    fill = {"Win": PatternFill("solid", fgColor="C8E6C9"), "Loss": PatternFill("solid", fgColor="FFCDD2"),
            "Open": PatternFill("solid", fgColor="BBDEFB"), "BE": PatternFill("solid", fgColor="FFE0B2")}
    records, auto_cache = [], {}
    for i in range(2, ws.max_row + 1):
        v = {h: ws.cell(i, col[h]).value for h in head}
        out = {c: None for c in NEW_COLS}
        t0 = int(v["Entry (UTC)"].replace(tzinfo=UTC).timestamp()) if v.get("Entry (UTC)") else None
        bars, tf = bars_for(v.get("Symbol"), t0)
        nd = config.decimals(v.get("Symbol"), v.get("Entry")) if v.get("Symbol") else 2
        if t0 and v.get("Symbol"):
            out.update({c: (round(x, nd + 1) if x is not None else None) for c, x in zip(ATR_COLS, indicators.atr_at(v["Symbol"], t0).values())})
            vw = indicators.vwap_at(v["Symbol"], t0)
            if vw and v.get("Entry") is not None and v.get("Stop") is not None and v.get("TP planned") is not None:
                s_ = 1 if v["Direction"] == "Long" else -1; e_ = v["Entry"]; dist = (e_ - vw[0]) * s_
                out["VWAP"] = round(vw[0], nd + 1); out["VWAP dist"] = round(dist, nd)
                out["VWAP side"] = "With" if (vw[2] - vw[0]) * s_ > 0 else "Against"
                out["VWAP in the way"] = "Yes" if 0 < -dist < abs(v["TP planned"] - e_) else "No"
                out["VWAP behind"] = "Yes" if 0 < dist < abs(e_ - v["Stop"]) else "No"
                out["VWAP 30m change"] = None if vw[1] is None else "With" if (vw[0] - vw[1]) * s_ > 0 else "Against"
                out["VWAP direction"], out["VWAP immediate slope"] = vwap_direction(indicators.vwap_recent(v["Symbol"], t0), s_)
        ok = bars and None not in (v.get("Entry"), v.get("Stop"), v.get("TP planned"), t0)
        if ok and v.get("Outcome") != "Not filled":
            long_ = v["Direction"] == "Long"; s = 1 if long_ else -1
            e, sl, tp, risk = v["Entry"], v["Stop"], v["TP planned"], abs(v["Entry"] - v["Stop"])
            res, mae, both, mfe = walk(bars, t0, long_, e, sl, tp)
            if res is not None:
                out["MAE pts"] = None if both else round(mae, nd)
                out["MFE pts"] = None if both else round(mfe, nd)
                out["1/2 stop"] = walk(bars, t0, long_, e, e - s * risk / 2, tp)[0]
                for t, c in zip(TARGETS, TARGET_COLS): out[c] = walk(bars, t0, long_, e, sl, e + s * t * risk)[0]
                out["BE at 55%"] = walk(bars, t0, long_, e, sl, tp, be_at=BE_AT)[0]
                out["1.5R BE at 55%"] = walk(bars, t0, long_, e, sl, e + s * 1.5 * risk, be_at=BE_AT)[0]
                out["1.5R BE at 1R"] = walk(bars, t0, long_, e, sl, e + s * 1.5 * risk, be_at=1 / 1.5)[0]
                a5 = out["ATR 5m"]
                if a5:
                    for (ks, kt), c in zip(ATR_RR, ATR_RR_COLS):
                        out[c] = walk(bars, t0, long_, e, e - s * ks * a5, e + s * kt * a5)[0]
        for c in NEW_COLS:
            cell = ws.cell(i, col[c], out[c])
            cell.fill = fill.get(out[c], PatternFill()); cell.alignment = Alignment(horizontal="center")
            if c in ATR_COLS or c in ("VWAP", "VWAP dist", "MAE pts", "MFE pts"): cell.number_format = "0." + "0" * max(1, nd) if c != "VWAP" else "0." + "0" * (nd + 1)
            good, bad = PatternFill("solid", fgColor="E8F5E9"), PatternFill("solid", fgColor="FFEBEE")
            if c in ("VWAP side", "VWAP 30m change", "VWAP immediate slope"): cell.fill = {"With": good, "Against": bad}.get(out[c], PatternFill())
            if c == "VWAP in the way": cell.fill = {"Yes": bad}.get(out[c], PatternFill())
            if c == "VWAP behind": cell.fill = {"Yes": good}.get(out[c], PatternFill())
            if c == "VWAP dist" and out[c] is not None and out[c] < 0: cell.fill = bad
        if "Auto confluence" in col and v.get("Symbol") not in auto_cache: auto_cache[v.get("Symbol")] = journal.auto_settings(wb, v.get("Symbol"))
        auto = auto_cache.get(v.get("Symbol")) if "Auto confluence" in col else None
        if auto and t0 and v.get("Symbol"):
            dist, unit, size, wanted = auto
            levels = {}
            for kind, n, name in wanted:
                if kind == "vwap": levels[name] = vw[0] if vw else None
                elif kind == "SMA": levels[name] = indicators.sma_at(v["Symbol"], t0, n)
                elif kind == "EMA": levels[name] = indicators.ema_at(v["Symbol"], t0, n)
                elif kind in ("previous day high", "previous day low"):
                    hl = indicators.prior_day_hl(v["Symbol"], t0)
                    levels[name] = (hl[0] if kind.endswith("high") else hl[1]) if hl else None
            text, slots = journal.auto_confluence(v, levels, dist, unit, size)
            ws.cell(i, col["Auto confluence"], text)
            for c, x in zip(journal.CONFS, slots): ws.cell(i, col[c], x); v[c] = x
            v["Auto confluence"] = text
        records.append((v, out))
    for c in NEW_COLS: ws.column_dimensions[get_column_letter(col[c])].width = 10
    ws.auto_filter.ref = ws.dimensions  # extract.py set it before these columns existed
    breakdown(wb, records)
    wb.save(path)
    try: add_chart_links(path)
    except PermissionError: pass
    return records


OFFICE_CONTAINER = Path.home() / "Library" / "Group Containers" / "UBF8T346G9.Office"


def jump_dir(page):
    """Where the one-line per-trade pages that jump to a trade's #anchor in the HTML gallery live.

    Excel for Mac is sandboxed: it asks permission for every file a link opens (and cannot be given a folder), and it
    strips both #anchors and ?queries from links. The one place it may open files without asking is Office's own group
    container, so on a Mac the jump pages go there (one folder per gallery) and the workbook links to them by absolute
    path. Elsewhere they sit beside the gallery and are linked relative to the workbook."""
    import hashlib, sys
    page = Path(page).resolve()
    if sys.platform == "darwin" and OFFICE_CONTAINER.exists():
        return OFFICE_CONTAINER / "tv-rr" / hashlib.sha1(str(page).encode()).hexdigest()[:12]
    return page.with_name(page.stem + "_files") / "trades"


def add_chart_links(path, page=None):
    """Fill the Chart column: for each trade, a link (relative to the workbook) to its section of the HTML gallery.
    Without `page`, look for the gallery beside the workbook or in gallery.notes_dir. Returns how many were linked."""
    import os, re
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    path = Path(path)
    lock = path.with_name("~$" + path.name)
    if lock.exists(): raise PermissionError(path)
    if page is None:
        dirs = [config.path(config.load()["gallery"].get("notes_dir")), path.parent]
        found = [p for d in dirs if d and d.exists() for p in d.glob("*- trade gallery.html")]
        if not found: return 0
        page = max(found, key=lambda p: p.stat().st_mtime)
    page = Path(page); tdir = jump_dir(page); in_container = OFFICE_CONTAINER in tdir.parents
    wb = load_workbook(path); ws = wb["Trades"]; head = [c.value for c in ws[1]]
    if "Chart" not in head:
        ws.cell(1, len(head) + 1, "Chart").font = Font(bold=True); head.append("Chart")
    col = {h: i + 1 for i, h in enumerate(head)}; n = 0
    for i in range(2, ws.max_row + 1):
        sym, did = ws.cell(i, col["Symbol"]).value, ws.cell(i, col["Drawing id"]).value
        cell = ws.cell(i, col["Chart"]); cell.value = None; cell.hyperlink = None
        if not sym or not did: continue
        target = tdir / ("t-" + re.sub(r"[^A-Za-z0-9_-]", "_", f"{sym}-{did}") + ".html")
        if not target.exists(): continue
        cell.value = "Open chart"
        cell.hyperlink = target.as_uri() if in_container else os.path.relpath(target, path.parent).replace(os.sep, "/")
        cell.font = Font(color="0563C1", underline="single"); n += 1
    wb.save(path)
    return n


def breakdown(wb, records):
    from openpyxl.styles import Font, PatternFill, Alignment
    if "Breakdown" in wb.sheetnames: del wb["Breakdown"]
    ws = wb.create_sheet("Breakdown")
    planned = lambda v, o: {"TP": "Win", "Stop": "Loss", "Both in one candle": "Loss", "Open": "Open"}.get(v.get("Outcome"))
    variants = [("Planned (your stop and TP)", planned, lambda v: (v.get("Reward pts"), v.get("Risk pts")))]
    variants.append(("1/2 stop (same TP)", lambda v, o: o["1/2 stop"],
                     lambda v: (v.get("Reward pts"), (v.get("Risk pts") or 0) / 2 or None)))
    variants.append((f"Stop to entry at {BE_AT:.0%} of TP", lambda v, o: o["BE at 55%"],
                     lambda v: (v.get("Reward pts"), v.get("Risk pts"))))
    r15 = lambda v: (1.5 * v["Risk pts"], v["Risk pts"]) if v.get("Risk pts") else (None, None)
    variants.append((f"1.5R target, stop to entry at {BE_AT:.0%}", lambda v, o: o["1.5R BE at 55%"], r15))
    variants.append(("1.5R target, stop to entry at +1R", lambda v, o: o["1.5R BE at 1R"], r15))
    for t, c in zip(TARGETS, TARGET_COLS):
        variants.append((f"{t:g}R target (your stop)", lambda v, o, c=c: o[c],
                         lambda v, t=t: (t * v["Risk pts"], v["Risk pts"]) if v.get("Risk pts") else (None, None)))
    for (ks, kt), c in zip(ATR_RR, ATR_RR_COLS):
        variants.append((f"Stop {ks:g}x / TP {kt:g}x 5m ATR", lambda v, o, c=c: o[c],
                         lambda v, ks=ks, kt=kt: (kt * v["ATR 5m"], ks * v["ATR 5m"]) if v.get("ATR 5m") else (None, None)))
    symbols = sorted({v.get("Symbol") for v, _ in records if v.get("Symbol")})
    groups = ["All", "Long", "Short"] + (symbols if len(symbols) > 1 else [])
    ws.append([f"Win/loss by variant — {len(records)} trades in the log, re-checked on the stored prices "
               f"(Planned uses your confirmed outcome where you set one)"])
    ws["A1"].font = Font(bold=True, size=13)
    sections = [(None, records)]
    if any("Decision" in v for v, _ in records):  # journal columns: also the trades that passed the filters
        unf = [(v, o) for v, o in records if v.get("Decision") != "Filtered"]
        sections = [(f"All trades ({len(records)}: taken, missed and filtered)", records),
                    (f"Unfiltered trades ({len(unf)}: taken and missed; Decision is not Filtered)", unf)]
    for title, recs in sections:
        variant_table(ws, title, recs, variants, groups)
    ws.append([])
    ws.append(["Win = TP reached before the stop; Loss = stop first; Open = neither yet; Break-even = the stop had been moved to the entry and price came back to it (0 pts; counted in the win % as not a win). Net R uses each variant's own risk "
               "(half the stop for 1/2 stop). Net pts adds price moves, so compare R when the log mixes instruments. MAE / MFE are the planned trade's worst move against and best move for the entry, averaged over the trades this variant won or lost; trades whose exit candle reached both levels have neither."])
    for col, w in zip("ABCDEFGHIJKLMN", (32, 14, 8, 8, 8, 8, 10, 8, 10, 8, 14, 14, 14, 14)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def variant_table(ws, title, records, variants, groups):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws.append([])
    if title:
        ws.append([title]); ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    hdr = ["Variant", "Group", "Trades", "Wins", "Losses", "Open", "Break-even", "Win %", "Net pts", "Net R", "Avg MAE winners", "Avg MAE losers", "Avg MFE winners", "Avg MFE losers"]
    ws.append(hdr)
    for c in ws[ws.max_row]: c.font = Font(bold=True); c.alignment = Alignment(horizontal="center", wrap_text=True)
    for name, res_of, rr_of in variants:
        for side in groups:
            rows = [(v, o) for v, o in records if side == "All" or side in (v.get("Direction"), v.get("Symbol"))]
            got = [(v, o, res_of(v, o)) for v, o in rows]
            got = [(v, o, r) for v, o, r in got if r in ("Win", "Loss", "Open", "BE")]
            w = sum(r == "Win" for *_, r in got); l = sum(r == "Loss" for *_, r in got); op = sum(r == "Open" for *_, r in got); be = sum(r == "BE" for *_, r in got)
            pts = R = 0.0
            for v, o, r in got:
                reward, risk = rr_of({**v, **o})
                if r == "Win" and reward: pts += reward; R += reward / risk if risk else 0
                if r == "Loss" and risk: pts -= risk; R -= 1
            pts = round(pts, 1 if abs(pts) >= 100 else 6)
            ex = lambda key, want: [o[key] for v, o, r in got if r == want and o[key] is not None]
            avg = lambda xs: (lambda m: round(m, 1 if abs(m) >= 100 else 2 if abs(m) >= 1 else 6))(statistics.mean(xs)) if xs else None
            ws.append([name, side, len(got), w, l, op, be, (w / (w + l + be)) if w + l + be else None, pts, round(R, 2),
                       avg(ex("MAE pts", "Win")), avg(ex("MAE pts", "Loss")), avg(ex("MFE pts", "Win")), avg(ex("MFE pts", "Loss"))])
            r_ = ws.max_row
            ws.cell(r_, 8).number_format = "0%"
            if side == "All":
                for c in ws[r_]: c.font = Font(bold=True)
                for c in ws[r_]: c.fill = PatternFill("solid", fgColor="F5F5F5")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--file"); a = ap.parse_args()
    path = config.workbook(a.file); recs = enrich(path)
    print(f"enriched {len(recs)} trades → {path} (columns {', '.join(NEW_COLS)}; sheet Breakdown)")
    for v, o in recs:
        print(f"  {v['Direction']:5} {v['Entry time']:%a %d %b %H:%M}  planned {v.get('Outcome') or '-':6} | "
              + " ".join(f"{c} {o[c] or '-'}" for c in NEW_COLS))
